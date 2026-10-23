"""Meaningful checks for supplied estimates, editable slide objects and preservation."""
import importlib.util,json,subprocess,sys,tempfile,unittest,zipfile
from pathlib import Path
R=Path(__file__).resolve().parents[1]

class VisualTests(unittest.TestCase):
    def run_script(self,name,*args):
        return subprocess.run([sys.executable,str(R/'skills/scripts'/name),*map(str,args)],capture_output=True,text=True,encoding='utf-8')

    def test_plot_rejects_invalid_interval_and_log_values(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);csv=p/'bad.csv'
            for body in ['a,2,3,4','a,0,-1,1','a,nan,1,2']:
                csv.write_text('label,estimate,lower,upper\n'+body+'\n',encoding='utf-8')
                out=self.run_script('plot_effects.py',csv,'--out',p/'bad','--scale','ratio')
                self.assertNotEqual(out.returncode,0)
                self.assertFalse((p/'bad.png').exists())

    def test_plot_exports_supplied_estimates(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);csv=p/'teaching-only.csv'
            csv.write_text('label,estimate,lower,upper\nTeaching A,1.2,1.0,1.5\nTeaching B,0.8,0.6,1.0\n',encoding='utf-8')
            out=self.run_script('plot_effects.py',csv,'--out',p/'teaching','--scale','ratio')
            self.assertEqual(out.returncode,0,out.stderr)
            for ext in ['svg','pdf','png']:self.assertGreater((p/f'teaching.{ext}').stat().st_size,100)
            report=json.loads((p/'teaching.json').read_text(encoding='utf-8'))
            self.assertEqual(report['rows'],2);self.assertFalse(report['model_fitted'])

    def test_deck_text_notes_and_bounds_audit(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);spec=p/'spec.json';deck=p/'deck.pptx'
            spec.write_text(json.dumps({'slides':[{'title':'Teaching example only','message':'A claim supported by supplied evidence','bullets':['First point','Second point'],'notes':'Speaker note retained','citation':'Teaching fixture; not a research result.'}]}),encoding='utf-8')
            out=self.run_script('build_editable_deck.py',spec,'--out',deck)
            self.assertEqual(out.returncode,0,out.stderr)
            prs=Presentation(deck);self.assertEqual(len(prs.slides),1)
            self.assertTrue(any(s.has_text_frame and s.text=='Teaching example only' for s in prs.slides[0].shapes))
            self.assertIn('Speaker note retained',prs.slides[0].notes_slide.notes_text_frame.text)
            clean=self.run_script('audit_slides.py',deck,'--out',p/'clean.json')
            self.assertEqual(clean.returncode,0,clean.stderr)
            self.assertEqual(json.loads((p/'clean.json').read_text())['visual_review'],'not_performed')
            prs.slides[0].shapes.add_textbox(Inches(14),Inches(1),Inches(1),Inches(1)).text='outside'
            prs.save(deck)
            bad=self.run_script('audit_slides.py',deck,'--out',p/'bad.json')
            self.assertEqual(bad.returncode,1,bad.stderr)
            self.assertTrue(json.loads((p/'bad.json').read_text())['errors'])

    def test_catalogue_addition_preserves_existing_records(self):
        rows=json.loads((R/'data/visual-skills.json').read_text(encoding='utf-8'))
        self.assertEqual(len(rows),12);self.assertEqual(len({x['id'] for x in rows}),12)
        old=json.loads(subprocess.check_output(['git','show','728fc01:data/studio-legacy.json'],cwd=R).decode('utf-8'))
        current=json.loads((R/'data/studio-legacy.json').read_text(encoding='utf-8'))
        self.assertEqual(current['SKILLS'][:len(old['SKILLS'])],old['SKILLS'])
        for key in old:
            if key!='SKILLS':self.assertEqual(current[key],old[key],key)
        self.assertEqual(len(current['SKILLS']),54)
        for r in rows:
            self.assertTrue((R/'skills'/(r['guide']+'.md')).is_file())
            self.assertTrue((R/'guides'/(r['guide']+'.html')).is_file())
        with zipfile.ZipFile(R/'data/figure-ppt-skills.zip') as z:
            self.assertEqual(len([n for n in z.namelist() if n.endswith('/SKILL.md')]),12)
            self.assertFalse(z.testzip())

if __name__=='__main__':unittest.main()
