"""Build a restrained editable PPTX draft from user-authored slide JSON.

JSON: {"font":"Arial", "slides":[{"title":"...", "message":"...",
"bullets":["..."], "figure":"relative/image.png", "citation":"...", "notes":"..."}]}
Only title is mandatory. Figure paths resolve beside the input JSON.
This script writes slides, not rendered layout certification.
"""
import argparse,json
from pathlib import Path

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('spec',type=Path);p.add_argument('--out',required=True,type=Path)
    a=p.parse_args();spec=json.loads(a.spec.read_text(encoding='utf-8-sig'))
    if not isinstance(spec,dict) or not isinstance(spec.get('slides'),list) or not spec['slides']:
        p.error('JSON requires a non-empty slides array')
    for i,row in enumerate(spec['slides'],1):
        if not isinstance(row,dict) or not isinstance(row.get('title'),str) or not row['title'].strip():
            p.error(f'Slide {i} needs a title')
        if len(row['title'])>85:p.error(f'Slide {i} title too long; shorten or split it')
        for key in ['message','citation','notes','figure']:
            if key in row and not isinstance(row[key],str):p.error(f'Slide {i} {key} must be text')
        bullets=row.get('bullets',[])
        if not isinstance(bullets,list) or not all(isinstance(b,str) for b in bullets):
            p.error(f'Slide {i} bullets must be a text array')
        if len(bullets)>5 or any(len(b)>110 for b in bullets):
            p.error(f'Slide {i} contains too much text for this draft layout; split it')
        if len(row.get('message',''))>160 or len(row.get('citation',''))>260:
            p.error(f'Slide {i} message or citation too long for this draft layout')
        if row.get('figure') and not (a.spec.parent/row['figure']).is_file():p.error(f'Slide {i} figure missing')
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from PIL import Image
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    font=spec.get('font','Arial')
    def textbox(slide,x,y,w,h,text,size,color='#253a34',bold=False):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        frame=shape.text_frame;frame.word_wrap=True
        frame.margin_left=frame.margin_right=Inches(.02)
        for i,line in enumerate(text.split('\n')):
            q=frame.paragraphs[0] if i==0 else frame.add_paragraph()
            q.text=line;q.font.name=font;q.font.size=Pt(size);q.font.bold=bold
            q.font.color.rgb=RGBColor.from_string(color.lstrip('#'));q.space_after=Pt(12)
        return shape
    for row in spec['slides']:
        s=prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(250,251,248)
        textbox(s,.65,.38,12,1.1,row['title'],29,bold=True)
        textbox(s,.65,1.55,12,.9,row.get('message',''),21)
        has_figure=bool(row.get('figure'))
        textbox(s,.65,2.65,4.25 if has_figure else 11.8,3.7,'\n'.join('• '+b for b in row.get('bullets',[])),21)
        if has_figure:
            path=a.spec.parent/row['figure']
            with Image.open(path) as im:ratio=im.width/im.height
            bw,bh=7.25,3.8;w=min(bw,bh*ratio);h=w/ratio
            s.shapes.add_picture(str(path),Inches(5.35+(bw-w)/2),Inches(2.55+(bh-h)/2),width=Inches(w),height=Inches(h))
        textbox(s,.65,6.65,12,.6,row.get('citation',''),11,color='#506159')
        if row.get('notes'):s.notes_slide.notes_text_frame.text=row['notes']
    a.out.parent.mkdir(parents=True,exist_ok=True);prs.save(a.out)
    print(f'Wrote {len(prs.slides)} editable draft slides. Render and review before use.')

if __name__=='__main__':main()
