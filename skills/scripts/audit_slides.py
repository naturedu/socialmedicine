"""Inspect PPTX bounds, notes and image coverage; not a visual overflow detector."""
import argparse,json
from pathlib import Path

def audit(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs=Presentation(path);slides=[];errors=[];warnings=[]
    def flatten(shapes):
        for shape in shapes:
            yield shape
            if shape.shape_type==MSO_SHAPE_TYPE.GROUP:yield from flatten(shape.shapes)
    for n,slide in enumerate(prs.slides,1):
        shapes=list(flatten(slide.shapes));pictures=[];texts=0
        for s in shapes:
            # Child coordinates can be transformed by group geometry; avoid false bounds claims.
            if s in slide.shapes:
                if s.left<0 or s.top<0 or s.left+s.width>prs.slide_width+9144 or s.top+s.height>prs.slide_height+9144:
                    errors.append({'slide':n,'shape':s.shape_id,'issue':'object_outside_canvas'})
            if getattr(s,'has_text_frame',False) and s.text.strip():texts+=1
            if s.shape_type==MSO_SHAPE_TYPE.PICTURE:
                coverage=s.width*s.height/(prs.slide_width*prs.slide_height)
                pictures.append({'shape':s.shape_id,'coverage':round(coverage,3)})
                if coverage>.85:warnings.append({'slide':n,'issue':'possible_full_slide_raster','coverage':round(coverage,3)})
        notes=slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ''
        if not notes:warnings.append({'slide':n,'issue':'no_speaker_notes'})
        slides.append({'slide':n,'shapes':len(shapes),'text_shapes':texts,'pictures':pictures,'has_notes':bool(notes)})
    if not slides:errors.append({'issue':'empty_deck'})
    return {'file':path.name,'slide_count':len(slides),'structure_passed':not errors,'errors':errors,
            'warnings':warnings,'slides':slides,'visual_review':'not_performed',
            'limits':['Does not verify text rendering, overlap, crop meaning or scientific correctness.',
                      'Nested group coordinate bounds are not certified.']}

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('pptx',type=Path)
    p.add_argument('--out',required=True,type=Path);a=p.parse_args();r=audit(a.pptx)
    a.out.parent.mkdir(parents=True,exist_ok=True);a.out.write_text(json.dumps(r,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(f"Slides: {r['slide_count']}; errors: {len(r['errors'])}; warnings: {len(r['warnings'])}; visual review required.")
    raise SystemExit(0 if r['structure_passed'] else 1)

if __name__=='__main__':main()
