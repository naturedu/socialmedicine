"""Extract public award metadata; no title selection or inferred award states here."""
from pathlib import Path
import re,json,hashlib,sys,warnings
from concurrent.futures import ProcessPoolExecutor
import pymupdf
warnings.filterwarnings('ignore',category=UserWarning)
R=Path(__file__).parent
def clean(x):return re.sub(r'\s+|_x000D_|_x000A_|_x0009_','',str(x if x is not None else ''),flags=re.I)
TITLE=['项目名称','课题名称','立项课题','中标选题研究方向']
def parse(meta):
 p=Path(meta['path']);out=Path(meta['cache'])/'parsed'/(meta['key']+'.json');out.parent.mkdir(exist_ok=True)
 if out.exists():return json.loads(out.read_text(encoding='utf-8'))['audit']
 rows=[];issues=[];headers=None;lasttitle='';lastserial=0
 def process(table,page=None,sheet=''):
  nonlocal headers,lasttitle,lastserial
  for rownum,line in enumerate(table,1):
   c=[clean(x) for x in line]
   if any(x in TITLE for x in c) and (any(x in ['序号','编号'] for x in c) or any(x in ['批准号','项目批准号'] for x in c)):
    headers=c
    if page is None:lasttitle=''
    continue
   if not c:continue
   merged_serial=False
   if re.fullmatch(r'\d+(?:\.0)?',c[0]):serial=int(float(c[0]));lastserial=serial
   elif meta.get('merged_titles') and headers and any(re.fullmatch(r'\d{2}(?:ZDA|&ZD)\d+',x) for x in c):serial=lastserial;merged_serial=True
   else:continue
   if not headers:issues.append(dict(serial=serial,page=page,sheet=sheet,error='missing header'));continue
   def get(*names):
    for name in names:
     if name in headers and headers.index(name)<len(c):return c[headers.index(name)]
    return ''
   title=get(*TITLE)
   if not title and meta.get('merged_titles'):title=lasttitle
   if not title:issues.append(dict(serial=serial,page=page,sheet=sheet,error='missing title'));continue
   lasttitle=title
   code=get('项目批准号','批准号','项目编号')
   organization=get('学校名称','工作单位','责任单位','项目单位')
   principal=get('申请人','姓名','负责人','首席专家','项目负责人','项目负责人姓名')
   if not organization:issues.append(dict(serial=serial,page=page,sheet=sheet,error='no organization',title=title));continue
   year=int('20'+code[:2]) if re.match(r'^\d{2}[A-Z&]',code) else (meta.get('year') or int(re.search(r'20\d\d',sheet)[0]))
   rawkind=get('项目类别')
   kind=meta.get('project_type')
   if not kind:
    kind={'A':'重点项目','B':'一般项目','C':'青年项目','X':'西部项目','F':'后期资助'}.get(code[2:3],rawkind or '年度项目')
   # Later-funding subcategories are not confused with annual A/B awards.
   if meta.get('program')=='后期资助':kind='后期资助'+(' · '+rawkind if rawkind else '')
   rows.append(dict(project_code=code,serial=serial,date=str(year),title=title,organization=organization,principal=principal,
     category=get('学科门类','学科','所在学科','涉及学科','学科分类','所属学科'),project_type=kind,program=meta.get('program','年度项目'),
     source_id=meta['key'],sheet=sheet,page=page,table_row=rownum,
     locator=('第 '+str(page)+' 页 · 序号 '+str(serial)) if page else (sheet+' · 序号 '+str(serial)+('（合并序号）' if merged_serial else '')+' · 表内行 '+str(rownum))))
 if p.suffix=='.pdf':
  for i,page in enumerate(pymupdf.open(p),1):
   for t in page.find_tables().tables:process(t.extract(),i)
 else:
  if p.suffix=='.xls':
   import xlrd
   for s in xlrd.open_workbook(p).sheets():
    if meta.get('sheets') and s.name not in meta['sheets']:continue
    headers=None;lasttitle='';process([s.row_values(i) for i in range(s.nrows)],sheet=s.name)
  else:
   import openpyxl
   for s in openpyxl.load_workbook(p,read_only=True,data_only=True).worksheets:
    if meta.get('sheets') and s.title not in meta['sheets']:continue
    headers=None;lasttitle='';process(list(s.iter_rows(max_col=12,values_only=True)),sheet=s.title)
 audit=dict(key=meta['key'],parsedRows=len(rows),sha256=hashlib.sha256(p.read_bytes()).hexdigest(),issues=issues)
 out.write_text(json.dumps(dict(audit=audit,rows=rows),ensure_ascii=False,indent=2),encoding='utf-8')
 print(meta['key'],len(rows),'rows',len(issues),'issues',flush=True)
 return audit

def run():
 import argparse
 from concurrent.futures import ThreadPoolExecutor
 ap=argparse.ArgumentParser(description=__doc__)
 ap.add_argument('--manifest',type=Path,default=Path(__file__).resolve().parents[1]/'data/nssfc-source-files.json')
 ap.add_argument('--cache',type=Path,required=True)
 ap.add_argument('--offline',action='store_true')
 args=ap.parse_args();cache=args.cache.resolve();cache.mkdir(parents=True,exist_ok=True)
 manifest=json.loads(args.manifest.read_text(encoding='utf-8'))
 def prepare(m):
  p=cache/m['filename']
  if not p.exists():
   if args.offline:raise FileNotFoundError(p)
   import requests
   from urllib.parse import urlsplit
   n=urlsplit(m['notice']);ref=n.scheme+'://'+n.netloc+'/'
   response=requests.get(m['url'],headers={'User-Agent':'Mozilla/5.0','Referer':ref},timeout=(8,35))
   response.raise_for_status();blob=response.content
   if hashlib.sha256(blob).hexdigest()!=m['sha256']:raise ValueError(m['key']+' source has changed; review before updating its hash')
   p.write_bytes(blob)
  if hashlib.sha256(p.read_bytes()).hexdigest()!=m['sha256']:raise ValueError(m['key']+' cached file hash mismatch')
  return {**m,'path':str(p),'cache':str(cache)}
 with ThreadPoolExecutor(max_workers=3) as pool:jobs=list(pool.map(prepare,manifest))
 with ProcessPoolExecutor(max_workers=3) as pool:audits=list(pool.map(parse,jobs))
 (cache/'parse-audit.json').write_text(json.dumps(audits,ensure_ascii=False,indent=2),encoding='utf-8')
if __name__=='__main__':run()
